import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from content_preliminary import add_preliminary_pages
from content_ch1_to_ch3 import add_chapter_1, add_chapter_2, add_chapter_3
from content_ch4_to_ch5 import add_chapter_4, add_chapter_5
from content_ch6_to_ch7 import add_chapter_6, add_chapter_7
from content_ch8_to_ch10 import add_chapter_8, add_chapter_9, add_chapter_10

def generate_full_docx():
    print("Initializing docx generation...")
    doc = docx.Document()

    # Configure Margins (Left 1.25", Right 1.0", Top 1.0", Bottom 1.0")
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Configure Default Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)
    style_normal.paragraph_format.space_before = Pt(0)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    img_dir = r"d:\Smart Waredrob\Document\extracted_images"

    # Add Preliminary Pages (Title, Declaration, Certificate, Acknowledgement, Abstract)
    add_preliminary_pages(doc, img_dir)

    # Add Chapters 1 to 10
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_chapter_5(doc, img_dir)
    add_chapter_6(doc)
    add_chapter_7(doc)
    add_chapter_8(doc)
    add_chapter_9(doc)
    add_chapter_10(doc)

    out_path = r"d:\Smart Waredrob\Document\Smart_Wardrobe_Mini_Project_Report.docx"
    doc.save(out_path)
    print(f"Document saved successfully to: {out_path}")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")

if __name__ == "__main__":
    generate_full_docx()
